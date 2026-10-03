# Тесты H (работа с шаблонами бегунков, настоящие .docx):
# импорт/скачивание/удаление/предпросмотр файлов в FILES_DIR/templates/,
# рендер из файла через docxtpl ({{ fio }} + {%tr%}-таблица по steps),
# фолбэк на текстовый body при отсутствии файла. Все ПДн вымышленные.
# Офлайн (нет docxtpl/qrcode) тест-рендер и рендер-тесты пропускаются
# через pytest.importorskip, структурная проверка работает всегда.

from __future__ import annotations

import importlib.util
import io
import os
import sys
import zipfile

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from app import docs as docs_module  # noqa: E402
from app.config import Settings, get_settings  # noqa: E402
from app.docs import generate_bypass  # noqa: E402
from app.main import app  # noqa: E402
from app.settings_routes import _templates_dir, get_settings_store  # noqa: E402

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _make_zip_docx(with_content_types: bool = True, with_document: bool = True) -> bytes:
    """Минимальный zip с частями .docx (для структурной проверки импорта)."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as archive:
        if with_content_types:
            archive.writestr(
                "[Content_Types].xml",
                "<?xml version='1.0'?><Types xmlns='http://schemas.openxmlformats.org/package/2006/content-types'/>",
            )
        if with_document:
            archive.writestr(
                "word/document.xml",
                "<?xml version='1.0'?><w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'/>",
            )
    return buf.getvalue()


def _make_uploadable_docx() -> bytes:
    """Валидный для импорта .docx: реальный (если есть python-docx) либо
    структурно корректный zip (офлайн: тест-рендер при импорте не выполняется)."""
    try:
        from docx import Document
    except ImportError:
        return _make_zip_docx()
    doc = Document()
    doc.add_paragraph("Сотрудник: {{ fio }}")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_template_docx_with_table() -> bytes:
    """Шаблон .docx для рендер-теста: {{ fio }} + {%tr%}-таблица по steps.

    Раскладка меток docxtpl: строка с `{%tr for %}` и строка с `{%tr endfor %}`
    — отдельные строки (в одной строке две `{%tr %}`-метки docxtpl не понимает,
    см. TEMPLATES.md), строка данных с `{{ step.* }}` — между ними.
    """
    from docx import Document

    doc = Document()
    doc.add_paragraph("Сотрудник: {{ fio }}")
    table = doc.add_table(rows=4, cols=3)
    head = table.rows[0].cells
    head[0].text = "№"
    head[1].text = "Владелец шага"
    head[2].text = "Статус"
    opener = table.rows[1].cells
    opener[0].text = "{%tr for step in steps %}"
    data_row = table.rows[2].cells
    data_row[0].text = "{{ step.order }}"
    data_row[1].text = "{{ step.owner }}"
    data_row[2].text = "{{ step.status }}"
    closer = table.rows[3].cells
    closer[0].text = "{%tr endfor %}"
    doc.add_paragraph("QR: {{ qr }}")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
def settings_override(tmp_path):
    """Тестовые группы + FILES_DIR в tmp_path (файлы шаблонов — сюда)."""
    settings = Settings(
        ALLOWED_AD_GROUPS="SED_HR,SED_ADMINS,SED_HR_ADMIN",
        ADMIN_GROUPS="SED_ADMINS",
        HR_GROUPS="SED_HR",
        HR_ADMIN_GROUPS="SED_HR_ADMIN",
        STEP_GROUP_PREFIX="SED_STEP_",
        FILES_DIR=str(tmp_path / "files"),
    )
    app.dependency_overrides[get_settings] = lambda: settings
    yield settings
    app.dependency_overrides.pop(get_settings, None)


class MemSettingsStore:
    """In-memory хранилище настроек (сид-формат), живёт весь тест
    (в отличие от autouse-подмены conftest, дающей пустой стор на каждый запрос)."""

    def __init__(self) -> None:
        self.values: dict[str, str] = {}

    def get(self, key: str) -> str | None:
        return self.values.get(key)

    def set(self, key: str, value: str) -> None:
        self.values[key] = value

    def get_many(self, keys) -> dict[str, str | None]:
        return {key: self.values.get(key) for key in keys}

    def set_many(self, values) -> None:
        self.values.update(values)


@pytest.fixture
def settings_store(settings_override):
    """Один in-memory стор на весь тест (PUT → GET из одного хранилища)."""
    store = MemSettingsStore()
    app.dependency_overrides[get_settings_store] = lambda: store
    yield store
    app.dependency_overrides.pop(get_settings_store, None)


def _upload(client, headers, content: bytes, filename: str = "blank.docx", previous: str | None = None):
    """POST /settings/doc-templates/files/upload (multipart)."""
    data = {"previous": previous} if previous is not None else {}
    return client.post(
        "/settings/doc-templates/files/upload",
        files={"file": (filename, content, DOCX_MIME)},
        data=data,
        headers=headers,
    )


# --- Валидация импорта (всегда, офлайн-путь) ---

def test_upload_rejects_non_zip(client, admin_headers, settings_override):
    """Простой текст вместо .docx — 400 (не zip), файл не сохраняется."""
    response = _upload(client, admin_headers, "это не zip-архив".encode("utf-8"))
    assert response.status_code == 400
    assert "zip" in response.json()["detail"].lower()
    assert not list(_templates_dir(settings_override.FILES_DIR).glob("*.docx"))


def test_upload_rejects_wrong_extension(client, admin_headers, settings_override):
    """Файл не с расширением .docx — 400 (до разбора содержимого)."""
    response = _upload(client, admin_headers, _make_zip_docx(), filename="blank.txt")
    assert response.status_code == 400
    assert ".docx" in response.json()["detail"]


def test_upload_rejects_zip_without_content_types(client, admin_headers, settings_override):
    """zip без [Content_Types].xml — 400."""
    response = _upload(client, admin_headers, _make_zip_docx(with_content_types=False))
    assert response.status_code == 400
    assert "[Content_Types].xml" in response.json()["detail"]


def test_upload_rejects_zip_without_document_xml(client, admin_headers, settings_override):
    """zip без word/document.xml — 400."""
    response = _upload(client, admin_headers, _make_zip_docx(with_document=False))
    assert response.status_code == 400
    assert "word/document.xml" in response.json()["detail"]


def test_upload_accepts_valid_docx_offline(client, admin_headers, settings_override):
    """Структурно валидный .docx — 200, файл в FILES_DIR/templates/ с uuid-именем."""
    content = _make_uploadable_docx()  # один раз: python-docx печатает в zip время сборки
    response = _upload(client, admin_headers, content)
    assert response.status_code == 200, response.text
    name = response.json()["name"]
    assert name.endswith(".docx")
    target = _templates_dir(settings_override.FILES_DIR) / name
    assert target.is_file()
    assert target.read_bytes() == content


def test_upload_previous_replaced(client, admin_headers, settings_override):
    """previous — старый файл удаляется после сохранения нового."""
    first = _upload(client, admin_headers, _make_uploadable_docx()).json()["name"]
    old_path = _templates_dir(settings_override.FILES_DIR) / first
    assert old_path.is_file()
    response = _upload(client, admin_headers, _make_uploadable_docx(), previous=first)
    assert response.status_code == 200, response.text
    second = response.json()["name"]
    assert second != first
    assert not old_path.exists()
    assert (_templates_dir(settings_override.FILES_DIR) / second).is_file()


def test_upload_rejects_broken_template(client, admin_headers, settings_override):
    """Битый шаблон (невалидный Jinja) — 400 (тест-рендер), файл не остаётся."""
    pytest.importorskip("docxtpl")
    from docx import Document

    doc = Document()
    doc.add_paragraph("Битый шаблон: {{ fio")
    buf = io.BytesIO()
    doc.save(buf)
    response = _upload(client, admin_headers, buf.getvalue(), filename="broken.docx")
    assert response.status_code == 400
    assert "не рендерится" in response.json()["detail"]
    assert not list(_templates_dir(settings_override.FILES_DIR).glob("*.docx"))


# --- Доступ (контент: админ + руководитель ОК) ---

def test_template_files_hr_forbidden(client, hr_headers, settings_override):
    """ОК (не контент-роль) — 403 на все операции с файлами."""
    content = _make_zip_docx()
    assert _upload(client, hr_headers, content).status_code == 403
    assert client.get("/settings/doc-templates/files/x.docx/download", headers=hr_headers).status_code == 403
    assert client.delete("/settings/doc-templates/files/x.docx", headers=hr_headers).status_code == 403
    assert client.post("/settings/doc-templates/files/x.docx/preview", headers=hr_headers).status_code == 403


def test_template_files_noauth_401(client, noauth_headers, settings_override):
    """Без логина — 401."""
    response = _upload(client, noauth_headers, _make_zip_docx())
    assert response.status_code == 401


def test_template_files_hr_admin_allowed(client, hr_admin_headers, settings_override):
    """Руководитель ОК (контент-роль) импортирует файл — 200."""
    response = _upload(client, hr_admin_headers, _make_uploadable_docx())
    assert response.status_code == 200, response.text
    assert response.json()["name"].endswith(".docx")


# --- Скачивание / удаление / предпросмотр ---

def test_template_files_download_roundtrip(client, admin_headers, settings_override):
    """Скачивание возвращает загруженные байты (attachment, имя — basename)."""
    content = _make_uploadable_docx()
    name = _upload(client, admin_headers, content).json()["name"]
    response = client.get(f"/settings/doc-templates/files/{name}/download", headers=admin_headers)
    assert response.status_code == 200
    assert response.content == content
    assert "attachment" in response.headers.get("content-disposition", "")
    assert name in response.headers.get("content-disposition", "")


def test_template_files_download_404(client, admin_headers, settings_override):
    """Несуществующий файл — 404; имя с обходом каталога тоже (basename)."""
    assert client.get("/settings/doc-templates/files/absent.docx/download", headers=admin_headers).status_code == 404
    assert client.get("/settings/doc-templates/files/../secret.docx/download", headers=admin_headers).status_code == 404


def test_template_files_delete(client, admin_headers, settings_override):
    """Удаление: файл пропадает, повтор — 404, ответ {"ok": true}."""
    name = _upload(client, admin_headers, _make_uploadable_docx()).json()["name"]
    response = client.delete(f"/settings/doc-templates/files/{name}", headers=admin_headers)
    assert response.status_code == 200
    assert response.json() == {"ok": True}
    assert not (_templates_dir(settings_override.FILES_DIR) / name).exists()
    assert client.delete(f"/settings/doc-templates/files/{name}", headers=admin_headers).status_code == 404
    assert client.delete("/settings/doc-templates/files/../secret.docx", headers=admin_headers).status_code == 404


def test_template_files_preview_offline_reason(client, admin_headers, settings_override):
    """Офлайн-предпросмотр — 200 {"generated": false, reason} (не 500)."""
    name = _upload(client, admin_headers, _make_uploadable_docx()).json()["name"]
    response = client.post(f"/settings/doc-templates/files/{name}/preview", headers=admin_headers)
    assert response.status_code == 200
    body = response.json()
    assert body["generated"] is False
    assert body["reason"]


def test_template_files_preview_404(client, admin_headers, settings_override):
    """Предпросмотр несуществующего файла — 404."""
    assert client.post("/settings/doc-templates/files/absent.docx/preview", headers=admin_headers).status_code == 404


# --- Поле file шаблона бегунка в настройках (GET/PUT /settings и /settings/content) ---

def test_settings_put_doc_template_file_persists(client, admin_headers, settings_override, settings_store):
    """Поле file проходит PUT /settings (exclude_unset) и возвращается GET."""
    payload = {
        "doc_templates": [
            {"service": "Служба вымышленного учета", "category": "линейный",
             "body": "Бегунок {{ fio }}", "file": "abc123.docx"}
        ]
    }
    response = client.put("/settings", json=payload, headers=admin_headers)
    assert response.status_code == 200, response.text
    assert response.json()["doc_templates"][0]["file"] == "abc123.docx"
    got = client.get("/settings", headers=admin_headers).json()
    assert got["doc_templates"][0]["file"] == "abc123.docx"


def test_settings_content_put_doc_template_file(client, hr_admin_headers, settings_override, settings_store):
    """Поле file проходит и через контент-эндпоинт (doc_templates — контент-ключ)."""
    response = client.put(
        "/settings/content",
        json={"doc_templates": [{"service": "Служба вымышленного учета", "category": "линейный",
                                 "body": "Бегунок {{ fio }}", "file": "abc123.docx"}]},
        headers=hr_admin_headers,
    )
    assert response.status_code == 200, response.text
    assert response.json()["doc_templates"][0]["file"] == "abc123.docx"


# --- Рендер из файла (docxtpl; офлайн — skip) ---

def test_render_from_file_substitutes_context(tmp_path, monkeypatch):
    """Рендер .docx-файла: {{ fio }} и {%tr%}-строки по steps подставлены."""
    pytest.importorskip("docxtpl")
    tpl_dir = tmp_path / "templates"
    tpl_dir.mkdir()
    tpl = tpl_dir / "blank.docx"
    tpl.write_bytes(_make_template_docx_with_table())
    # QR — заглушка (stdlib PNG): тест не зависит от qrcode, только от docxtpl.
    monkeypatch.setattr(docs_module, "_render_qr_png", lambda url: docs_module.qr_stub_png(url))
    context = {
        "fio": "Иванов Иван Иванович",
        "steps": [
            {"order": 1, "owner": "Группа первая", "status": "На согласовании"},
            {"order": 2, "owner": "Группа вторая", "status": "Согласовано"},
        ],
    }
    docx_bytes = docs_module._render_docx_from_file(tpl, context, "https://x/REQ-0001")
    from docx import Document

    rendered = Document(io.BytesIO(docx_bytes))
    paragraphs = [p.text for p in rendered.paragraphs]
    assert any("Иванов Иван Иванович" in text for text in paragraphs)
    cells = [c.text for t in rendered.tables for r in t.rows for c in r.cells]
    assert "Группа первая" in cells
    assert "Группа вторая" in cells
    assert "{{" not in " ".join(cells)


# --- Выбор пути рендера в generate_bypass (офлайн-безопасно) ---

def test_generate_bypass_uses_file_render_when_file_exists(tmp_path, monkeypatch):
    """template_file задан и файл есть — рендер из файла (не фолбэк)."""
    tpl_dir = tmp_path / "templates"
    tpl_dir.mkdir()
    tpl = tpl_dir / "blank.docx"
    tpl.write_bytes(b"docx-bytes")
    calls = {}

    def fake_from_file(path, context, url):
        calls["from_file"] = (path, url)
        return b"docx-from-file"

    def fake_stand(body, context):
        raise AssertionError("фолбэк на текстовый body не должен вызываться при наличии файла")

    monkeypatch.setattr(docs_module, "_render_docx_from_file", fake_from_file)
    monkeypatch.setattr(docs_module, "_render_docx_stand", fake_stand)
    monkeypatch.setattr(docs_module, "_render_qr_png", lambda url: b"qr")
    monkeypatch.setattr(docs_module, "_soffice_binary", lambda: None)
    result = generate_bypass(
        request_id="REQ-0001", version="v1", template_body="Текст {{ fio }}",
        context={"fio": "Иванов"}, base_url="https://x",
        files_dir=str(tmp_path), template_file="blank.docx",
    )
    assert calls["from_file"][0] == tpl
    assert calls["from_file"][1] == "https://x/requests/REQ-0001"
    assert result.generated is False  # нет soffice — PDF не создан


def test_generate_bypass_fallback_when_file_missing(tmp_path, monkeypatch):
    """template_file задан, но файла нет — фолбэк на текстовый body."""
    calls = []

    def fake_stand(body, context):
        calls.append(body)
        return b"docx-from-text"

    def fake_from_file(path, context, url):
        raise AssertionError("рендер из файла не должен вызываться без файла")

    monkeypatch.setattr(docs_module, "_render_docx_stand", fake_stand)
    monkeypatch.setattr(docs_module, "_render_docx_from_file", fake_from_file)
    monkeypatch.setattr(docs_module, "_render_qr_png", lambda url: b"qr")
    monkeypatch.setattr(docs_module, "_soffice_binary", lambda: None)
    generate_bypass(
        request_id="REQ-0001", version="v1", template_body="Бегунок: {{ fio }}",
        context={"fio": "Иванов"}, base_url="https://x",
        files_dir=str(tmp_path), template_file="absent.docx",
    )
    assert calls == ["Бегунок: {{ fio }}"]


def test_generate_bypass_no_template_file_uses_text(tmp_path, monkeypatch):
    """template_file не задан — текстовый рендер (старый контракт не меняется)."""
    calls = []

    def fake_stand(body, context):
        calls.append(body)
        return b"docx-from-text"

    monkeypatch.setattr(docs_module, "_render_docx_stand", fake_stand)
    monkeypatch.setattr(docs_module, "_render_qr_png", lambda url: b"qr")
    monkeypatch.setattr(docs_module, "_soffice_binary", lambda: None)
    generate_bypass(
        request_id="REQ-0001", version="v1", template_body="Текст {{ fio }}",
        context={"fio": "Иванов"}, base_url="https://x", files_dir=str(tmp_path),
    )
    assert calls == ["Текст {{ fio }}"]


@pytest.mark.skipif(
    importlib.util.find_spec("docxtpl") is not None,
    reason="ветка офлайна: docxtpl в окружении есть — ImportError не возникнет",
)
def test_generate_bypass_file_offline_reason(tmp_path, monkeypatch):
    """Файл есть, но нет docxtpl — ImportError → generated=False с причиной «офлайн»."""
    tpl_dir = tmp_path / "templates"
    tpl_dir.mkdir()
    (tpl_dir / "blank.docx").write_bytes(b"docx-bytes")
    monkeypatch.setattr(docs_module, "_render_qr_png", lambda url: b"qr")
    result = generate_bypass(
        request_id="REQ-0001", version="v1", template_body="Текст",
        context={}, base_url="https://x", files_dir=str(tmp_path), template_file="blank.docx",
    )
    assert result.generated is False
    assert result.reason.startswith("офлайн")
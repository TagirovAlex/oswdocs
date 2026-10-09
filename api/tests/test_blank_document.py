# Тесты печати бланка ИЗ ДАННЫХ (волна «Справочник бланков»):
# генератор docs.build_blank_document (python-docx, два пресета макета из снимка
# blank_layout), санирование текстов этапа (HTML из визуального редактора),
# QR на заявку, ответственные (ФИО персональных исполнителей / группы AD),
# generate_bypass без файлов-шаблонов и ручка печати (макет в blank_kind,
# пустые шаги — 422). Все ПДн вымышленные.

from __future__ import annotations

import base64
import inspect
import io
import os
import sys
import zipfile
from datetime import date, datetime, timedelta
from types import SimpleNamespace

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from app import docs as docs_module  # noqa: E402
from app.config import Settings, get_settings  # noqa: E402
from app.docs import (  # noqa: E402
    DEFAULT_BLANK_LAYOUT,
    LAYOUT_PRESETS,
    BypassResult,
    blank_placeholders,
    build_blank_document,
    generate_bypass,
    render_blank_text,
    resolve_blank_layout,
)
from app.main import app  # noqa: E402
from app.requests import _Request, _Step, _utcnow, get_memory_requests_store  # noqa: E402
from app.requests import (  # noqa: E402
    IN_APPROVAL,
    STEP_PENDING,
    get_requests_store,
)
from app.settings_routes import get_settings_store  # noqa: E402

TEST_ALLOWED = "SED_HR,SED_ADMINS"
TEST_ADMINS = "SED_ADMINS"
TEST_HR = "SED_HR"
TEST_STEP_PREFIX = "SED_STEP_"
BASE_URL = "https://sed-mock.local"

FAKE_ENTERPRISE_CODE = "ВЫМЫШЛЕННОЕ-ПРЕДПРИЯТИЕ"
FAKE_ENTERPRISE_NAME = "Вымышленное предприятие"
FAKE_SERVICE = "Служба вымышленного учета"
FAKE_POSITION = "Старший вымышленный кассир"
FAKE_FIO = "Вымышленный Сотрудник Полный"
FAKE_TAB_NUM = "В-0001"
FAKE_BLANK = "Бланк вымышленного увольнения"
FAKE_STAGE = "Бухгалтерия вымышленная"
FAKE_ASSIGNEE_FIO = "Бухгалтер Вымышленный"
FAKE_GROUP = "SED_STEP_BUH"
FAKE_SUBJECT = "Вымышленная тема заявки"
FAKE_CONTENT = "Вымышленное содержание заявки"
FAKE_MANAGER_FIO = "Руководитель Вымышленный"
FAKE_DATE = "2026-10-01"
FAKE_DISMISSAL_DATE = "2026-10-15"
QR_URL = BASE_URL + "/requests/REQ-0001"

# Шапка и подвал бланка из снимка заявки: HTML редактора и список строк с
# плейсхолдерами контракта печати.
FAKE_HEADER = (
    "<p><b>Акт вымышленного увольнения</b> № {tab_num}</p>"
    "<p>Сотрудник: {fio}, {position}</p>"
    "<p>Служба: {department}; предприятие: {enterprise}</p>"
    "<p>Печать: {date}; увольнение: {dismissal_date}</p>"
    "<p>{subject} — {content}</p>"
    "<p>Бланк {blank_name}, {steps}, руководитель {manager}</p>"
)
FAKE_FOOTER = [
    "Подпись сотрудника ____________________",
    "Дата: {date}, ФИО: {fio}",
]


def _b64(value: str) -> str:
    return base64.b64encode(value.encode("utf-8")).decode("ascii")


def _hr_headers() -> dict:
    """Заголовки сотрудника ОК (роль печати)."""
    return {
        "X-Mock-Sam": "ok.vymyshlennaya",
        "X-Mock-Fio": _b64("Вымышленный Пользователь Тестовый"),
        "X-Mock-Mail": _b64("ok.vymyshlennaya@example.com"),
        "X-Mock-Department": _b64("Вымышленный отдел"),
        "X-Mock-Title": _b64("Вымышленная должность"),
        "X-Mock-Groups": TEST_HR,
    }


def _context(**overrides) -> dict:
    """Контекст бланка из данных заявки (как build_bypass_context documents.py)."""
    context = {
        "request_id": "REQ-0001",
        "fio": FAKE_FIO,
        "position": FAKE_POSITION,
        "department": FAKE_SERVICE,
        "tab_num": FAKE_TAB_NUM,
        "enterprise": FAKE_ENTERPRISE_NAME,
        "blank_name": FAKE_BLANK,
        "blank_version": 2,
        "blank_layout": "office",
        "qr_url": QR_URL,
        "subject": FAKE_SUBJECT,
        "content": FAKE_CONTENT,
        "date": FAKE_DATE,
        "dismissal_date": FAKE_DISMISSAL_DATE,
        "manager": FAKE_MANAGER_FIO,
        "steps": [
            {
                "order": 1,
                "owner": FAKE_GROUP,
                "title": FAKE_STAGE,
                "stage_lines": ["Проверить расчёты", "Подписать акт"],
                "fio": FAKE_ASSIGNEE_FIO,
                "status": "Согласовано",
                "done_at": "05.10.2026 12:00",
            },
            {
                "order": 2,
                "owner": "SED_STEP_HR",
                "title": "Отдел кадров вымышленный",
                "stage_lines": ["Оформить приказ"],
                "fio": "",
                "status": STEP_PENDING,
                "done_at": "",
            },
        ],
    }
    context.update(overrides)
    return context


def _docx_text(blob: bytes) -> str:
    """Текст собранного документа (XML как есть — тест ищет вхождения)."""
    with zipfile.ZipFile(io.BytesIO(blob)) as archive:
        return archive.read("word/document.xml").decode("utf-8")


def _docx_names(blob: bytes) -> list[str]:
    with zipfile.ZipFile(io.BytesIO(blob)) as archive:
        return archive.namelist()


def _opened(blob: bytes):
    from docx import Document

    return Document(io.BytesIO(blob))


def _steps_rows(blob: bytes) -> list[list[str]]:
    """Строки таблицы шагов (последняя таблица документа)."""
    table = _opened(blob).tables[-1]
    return [[cell.text for cell in row.cells] for row in table.rows]


def _body_paragraphs(blob: bytes) -> list:
    """Абзацы тела документа по порядку (таблицы в них не входят)."""
    return _opened(blob).paragraphs


def _body_blocks(blob: bytes) -> list[str]:
    """Типы блоков тела документа по порядку: p — абзац, tbl — таблица."""
    body = _opened(blob).element.body
    return [child.tag.split("}")[-1] for child in body]


def _request(**overrides) -> _Request:
    """Заявка со снимком бланка и одним шагом (печатная форма)."""
    fields = dict(
        id="REQ-0001",
        status=IN_APPROVAL,
        route_origin="auto",
        enterprise=FAKE_ENTERPRISE_CODE,
        fio=FAKE_FIO,
        tab_num=FAKE_TAB_NUM,
        department=FAKE_SERVICE,
        position=FAKE_POSITION,
        created_by="ok.vymyshlennaya",
        blank_id=7,
        blank_name=FAKE_BLANK,
        blank_version=2,
        blank_layout="line",
        steps=[
            _Step(
                order=1,
                owner_group=FAKE_GROUP,
                assignee="step.buhgalter",
                status=STEP_PENDING,
                expires_at=_utcnow() + timedelta(days=3),
                stage_title=FAKE_STAGE,
                stage_lines=["Проверить расчёты"],
            )
        ],
    )
    fields.update(overrides)
    return _Request(**fields)


@pytest.fixture(autouse=True)
def settings_override():
    settings = Settings(
        ALLOWED_AD_GROUPS=TEST_ALLOWED,
        ADMIN_GROUPS=TEST_ADMINS,
        HR_GROUPS=TEST_HR,
        STEP_GROUP_PREFIX=TEST_STEP_PREFIX,
        APP_BASE_URL=BASE_URL,
    )
    app.dependency_overrides[get_settings] = lambda: settings
    yield settings
    app.dependency_overrides.pop(get_settings, None)


class MemSettingsStore:
    """Хранилище настроек в сид-формате (значения — строки JSON)."""

    def __init__(self, values: dict | None = None) -> None:
        self.values = dict(values or {})

    def get(self, key):
        return self.values.get(key)


@pytest.fixture(autouse=True)
def settings_store():
    """doc_templates заполнен намеренно: печать его читать не должна."""
    store = MemSettingsStore(
        {
            "doc_templates": (
                '[{"service": "%s", "category": "линейный", '
                '"body": "ШАБЛОН-НЕ-ИСПОЛЬЗУЕТСЯ", "file": "blank.docx"}]'
                % FAKE_SERVICE
            ),
            "enterprises": (
                '[{"code": "%s", "name": "%s"}]'
                % (FAKE_ENTERPRISE_CODE, FAKE_ENTERPRISE_NAME)
            ),
        }
    )
    app.dependency_overrides[get_settings_store] = lambda: store
    yield store
    app.dependency_overrides.pop(get_settings_store, None)


@pytest.fixture
def requests_store():
    store = get_memory_requests_store()
    app.dependency_overrides[get_requests_store] = lambda: store
    store.reset()
    yield store
    store.reset()
    app.dependency_overrides.pop(get_requests_store, None)


def _fake_result(pdf_path: str) -> BypassResult:
    return BypassResult(
        generated=True,
        docx_path=pdf_path.replace(".pdf", ".docx"),
        pdf_path=pdf_path,
        qr_payload=QR_URL,
    )


# ---------------------------------------------------------------------------
# Выбор макета: единственная точка resolve_blank_layout + два пресета
# ---------------------------------------------------------------------------

def test_resolve_blank_layout_office_by_default():
    """Пустой/неизвестный снимок макета — office (дефолт миграции 0012)."""
    assert DEFAULT_BLANK_LAYOUT == "office"
    assert resolve_blank_layout(None) == "office"
    assert resolve_blank_layout("") == "office"
    assert resolve_blank_layout("  ") == "office"
    assert resolve_blank_layout("неизвестный макет") == "office"
    assert resolve_blank_layout("линейный") == "office"


def test_resolve_blank_layout_line_ignores_case_and_spaces():
    """line из снимка — макет line; регистр и пробелы не важны."""
    assert resolve_blank_layout("line") == "line"
    assert resolve_blank_layout(" LINE ") == "line"
    assert resolve_blank_layout("Line") == "line"


def test_layout_presets_cover_contract_layouts():
    """Пресеты ровно office/line, дефолт среди них, доли колонок дают 100%."""
    assert set(LAYOUT_PRESETS) == {"office", "line"}
    assert DEFAULT_BLANK_LAYOUT in LAYOUT_PRESETS
    for name, preset in LAYOUT_PRESETS.items():
        shares = preset["step_widths"]
        assert len(shares) == 4, name
        assert abs(sum(shares) - 1.0) < 0.001, name


# ---------------------------------------------------------------------------
# Генератор: шапка, таблица шагов, подвал
# ---------------------------------------------------------------------------

def test_build_blank_document_is_docx_with_header_table_and_footer():
    """DOCX из данных: название бланка, поля сотрудника, таблица шагов, подвал."""
    blob = build_blank_document(_context())
    assert blob[:2] == b"PK"
    text = _docx_text(blob)
    for expected in (
        FAKE_BLANK,      # название бланка из снимка
        FAKE_FIO,        # ФИО
        FAKE_POSITION,   # должность
        FAKE_SERVICE,    # служба
        FAKE_TAB_NUM,    # табельный номер
        FAKE_ENTERPRISE_NAME,  # предприятие
    ):
        assert expected in text, expected
    rows = _steps_rows(blob)
    assert rows[0] == ["№", "Этап", "Ответственный", "Отметка/дата"]
    assert rows[1][0] == "1"
    assert FAKE_STAGE in rows[1][1]      # название этапа
    assert "Проверить расчёты" in rows[1][1]  # его пункты
    assert rows[1][2] == FAKE_ASSIGNEE_FIO  # ФИО персонального исполнителя
    assert "Согласовано" in rows[1][3] and "05.10.2026 12:00" in rows[1][3]
    assert rows[2][2] == "SED_STEP_HR"   # группа AD — печатаем как есть
    footer = _opened(blob).sections[0].footer.paragraphs[0].text
    assert "REQ-0001" in footer and FAKE_BLANK in footer and "v2" in footer


def test_build_blank_document_step_mark_without_date_has_paper_line():
    """Шаг без даты: статус и место под росчерк (печать от руки)."""
    rows = _steps_rows(build_blank_document(_context()))
    marks = rows[2][3].splitlines()
    assert marks[0] == STEP_PENDING
    assert set(marks[1]) == {"_"}


def test_build_blank_document_layout_comes_from_snapshot():
    """Макет берётся из снимка blank_layout, а не из службы/настроек."""
    office = _opened(build_blank_document(_context(blank_layout="office")))
    line = _opened(build_blank_document(_context(blank_layout="line")))
    # Линейный пресет компактнее офисного (поля страницы и кегль меньше).
    assert line.sections[0].left_margin < office.sections[0].left_margin
    assert line.styles["Normal"].font.size < office.styles["Normal"].font.size
    # Без макета в снимке — office (дефолт), а не иной пресет.
    default = _opened(build_blank_document(_context(blank_layout="")))
    assert default.sections[0].left_margin == office.sections[0].left_margin


def test_build_blank_document_qr_image_and_caption():
    """QR подставлен картинкой (qr_url контекста), рядом читается ссылка."""
    blob = build_blank_document(_context())
    assert any(name.startswith("word/media/") for name in _docx_names(blob))
    assert QR_URL in _docx_text(blob)


def test_build_blank_document_without_qr_url_still_builds():
    """Без qr_url (его кладёт generate_bypass) документ тоже собирается."""
    context = _context()
    context.pop("qr_url")
    blob = build_blank_document(context)
    assert not any(name.startswith("word/media/") for name in _docx_names(blob))


def test_build_blank_document_drops_forbidden_context_fields():
    """ПДн-поля в бланок не попадают (чистит sanitize_context)."""
    blob = build_blank_document(_context(vacation_balance="ДВАДЦАТЬ ВОСЕМЬ ДНЕЙ"))
    assert "ДВАДЦАТЬ" not in _docx_text(blob)


def test_build_blank_document_title_falls_back_to_doc_type():
    """Без названия бланка печатаем вид документа (doc_type_code)."""
    context = _context()
    context["blank_name"] = ""
    context["doc_type_code"] = "УВОЛЬНЕНИЕ"
    assert "УВОЛЬНЕНИЕ" in _docx_text(build_blank_document(context))


# ---------------------------------------------------------------------------
# Санирование текстов этапа (HTML из визуального редактора)
# ---------------------------------------------------------------------------

def test_build_blank_document_keeps_only_allowed_tags():
    """В печать идут разрешённые теги; служебные и прочие — без разметки."""
    context = _context()
    context["steps"] = [
        {
            "order": 1,
            "owner": FAKE_GROUP,
            "title": "Этап <b>итог</b>",
            "stage_lines": [
                "Проверить <b>расчёты</b> и <i>долги</i>",
                "Подпись <script>alert(1)</script>",
                "Проверить <img src=x onerror=alert(2)> табель",
                "Сумма 5 < 7 и 8 > 3",
            ],
            "fio": "",
            "status": STEP_PENDING,
            "done_at": "",
        }
    ]
    blob = build_blank_document(context)
    text = _docx_text(blob)
    stage_text = _steps_rows(blob)[1][1]
    # Текст печатается, форматирование разрешённых тегов — работает.
    assert "Проверить расчёты и долги" in stage_text
    assert "Подпись" in stage_text
    assert "Проверить  табель" in stage_text
    # Содержимое служебных тегов и атрибуты прочих — на бланок не попадают.
    assert "alert" not in text
    assert "onerror" not in text
    assert "<img" not in text
    # Голые спецсимволы остаются текстом (Word экранирует их сам).
    assert "Сумма 5 < 7 и 8 > 3" in stage_text
    stage_paragraphs = _opened(blob).tables[-1].rows[1].cells[1].paragraphs
    bolded = [run.text for paragraph in stage_paragraphs for run in paragraph.runs if run.bold]
    italiced = [run.text for paragraph in stage_paragraphs for run in paragraph.runs if run.italic]
    assert "расчёты" in bolded
    assert "долги" in italiced


def test_build_blank_document_breaks_line_on_allowed_tag():
    """Разрешённый перенос строки (br) печатается переносом, а не тегом."""
    context = _context()
    context["steps"] = [
        {
            "order": 1,
            "owner": FAKE_GROUP,
            "title": "Этап",
            "stage_lines": ["Первый пункт<br>второй пункт"],
            "fio": "",
            "status": STEP_PENDING,
            "done_at": "",
        }
    ]
    assert _steps_rows(build_blank_document(context))[1][1].splitlines() == [
        "Этап",
        "Первый пункт",
        "второй пункт",
    ]


# ---------------------------------------------------------------------------
# Шапка и подвал бланка из снимка заявки + плейсхолдеры печати
# ---------------------------------------------------------------------------

def test_build_blank_document_header_html_comes_from_snapshot():
    """Шапка в снимке печатается как шапка бланка; сетку полей она заменяет."""
    blob = build_blank_document(_context(blank_header_html=FAKE_HEADER))
    paragraphs = [paragraph.text for paragraph in _body_paragraphs(blob)]
    # Заголовок бланка печатается как было, пользовательская шапка — под ним.
    assert paragraphs[0] == FAKE_BLANK
    assert paragraphs[1] == "Акт вымышленного увольнения № " + FAKE_TAB_NUM
    assert paragraphs[2] == f"Сотрудник: {FAKE_FIO}, {FAKE_POSITION}"
    assert f"Бланк {FAKE_BLANK}, количество шагов: 2, руководитель {FAKE_MANAGER_FIO}" in paragraphs
    # Шапка заменяет сетку полей: подписей полей на бланке нет, таблица шагов одна.
    assert "Табельный номер" not in _docx_text(blob)
    assert len(_opened(blob).tables) == 1


def test_build_blank_document_header_html_keeps_allowed_formatting():
    """Начертания разрешённых тегов шапки работают, как у текста этапа."""
    blob = build_blank_document(
        _context(blank_header_html="<p>Акт <b>срочный</b> и <i>важный</i></p>")
    )
    paragraphs = _opened(blob).paragraphs
    runs = [run for run in paragraphs[1].runs]
    bold = [run.text for run in runs if run.bold]
    italic = [run.text for run in runs if run.italic]
    assert bold == ["срочный"]
    assert italic == ["важный"]


def test_build_blank_document_header_html_is_sanitized():
    """Скрипты, iframe, обработчики событий и внешние ресурсы на бланок не идут."""
    blob = build_blank_document(
        _context(
            blank_header_html=(
                "<p>Акт <b>вымышленный</b> по заявке {fio}</p>"
                "<p><script>alert(1)</script></p>"
                "<p><iframe src=\"https://вымышленный-ресурс/кадр\"></iframe></p>"
                "<p><img src=\"https://вымышленный-ресурс/картинка.png\" "
                "onerror=\"alert(2)\"></p>"
            )
        )
    )
    text = _docx_text(blob)
    assert "alert" not in text
    assert "onerror" not in text
    assert "вымышленный-ресурс" not in text
    assert "<img" not in text
    paragraphs = [paragraph.text for paragraph in _body_paragraphs(blob)]
    assert paragraphs[1] == f"Акт вымышленный по заявке {FAKE_FIO}"


def test_build_blank_document_header_placeholder_value_stays_plain_text():
    """Значение {content}/{subject} печатается текстом: разметка из значения
    не разрывает абзац шапки и не даёт начертания.

    Порядок печати: в HTML-шапку значение подставляется ЭКРАНИРОВАННЫМ, поэтому
    </p> из свободного текста заявки не добавляет абзац, а <b>/<script> не
    разбираются как теги бланка (структура документа от значения не зависит)."""
    header = (
        "<p><b>Акт вымышленного увольнения</b></p>"
        "<p>Содержание: {content}</p>"
        "<p>Тема: {subject}</p>"
    )
    markup = "</p><b>жирно</b> и <script>alert(1)</script>"
    blob = build_blank_document(
        _context(
            subject=f"тема {markup}",
            content=f"прощайте {markup}",
            blank_header_html=header,
        )
    )
    # Шапка из разметки без значений: столько же абзацев — значение ни одного
    # не добавило (его </p> не разорвал абзацы шапки).
    plain = build_blank_document(
        _context(subject="тема", content="прощайте", blank_header_html=header)
    )
    paragraphs = _body_paragraphs(blob)
    assert len(paragraphs) == len(_body_paragraphs(plain))
    # Значение напечатано символами текста: разметка из него осталась текстом.
    assert [paragraph.text for paragraph in paragraphs[1:4]] == [
        "Акт вымышленного увольнения",
        f"Содержание: прощайте {markup}",
        f"Тема: тема {markup}",
    ]
    # Начертание осталось только у разметки самой шапки, значение его не дало.
    bold = [run.text for paragraph in paragraphs for run in paragraph.runs if run.bold]
    assert bold == [FAKE_BLANK, "Акт вымышленного увольнения"]


def test_build_blank_document_placeholder_markup_does_not_break_print():
    """Разметка в значении плейсхолдера не роняет печать ни в одном месте бланка.

    Шапка (HTML) — значение экранируется, подвал — обычная подстановка в текст;
    в обоих случаях документ собирается, значение печатается символами, а шапка,
    таблица шагов и подвал на месте."""
    markup = "</p><b>жирно</b> <script>alert(1)</script>"
    context = _context(
        subject=markup,
        content=markup,
        blank_header_html="<p>Содержание: {content}</p>",
        blank_footer_lines=["Подпись {subject}", "Подпись {не_существует}"],
    )
    context["steps"] = [
        dict(context["steps"][0], title="Этап {content}", stage_lines=["Пункт {subject}"])
    ]
    blob = build_blank_document(context)
    paragraphs = [paragraph.text for paragraph in _body_paragraphs(blob)]
    assert paragraphs[1] == f"Содержание: {markup}"
    # Подвал печатается целиком, неизвестный плейсхолдер остаётся текстом.
    assert paragraphs[-2:] == [f"Подпись {markup}", "Подпись {не_существует}"]
    # Содержащее значение из {content} печатается в тексте шага, документ цел.
    assert len(_opened(blob).tables) == 1
    assert _docx_names(blob)[0] == "[Content_Types].xml"


def test_build_blank_document_without_header_html_keeps_field_grid():
    """Шапки в снимке нет — прежнее поведение: офисный пресет печатает сетку полей."""
    grid = _opened(build_blank_document(_context())).tables[0]
    assert [row.cells[0].text.strip() for row in grid.rows] == [
        "ФИО:", "Должность:", "Служба:", "Табельный номер:", "Предприятие:",
    ]
    assert grid.cell(0, 1).text == FAKE_FIO
    # Линейный пресет — те же поля списком.
    line = _body_paragraphs(build_blank_document(_context(blank_layout="line")))
    assert f"Табельный номер: {FAKE_TAB_NUM}" in [
        paragraph.text for paragraph in line
    ]


def test_build_blank_document_footer_lines_come_after_steps_table():
    """Строки подвала печатаются после таблицы шагов, по порядку и по центру."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    blob = build_blank_document(
        _context(blank_footer_lines=list(FAKE_FOOTER), date=None, fio=FAKE_FIO)
    )
    # Последние блоки тела — абзацы подвала, таблица шагов перед ними.
    blocks = _body_blocks(blob)
    assert blocks[-1] == "sectPr"
    assert blocks[-3:-1] == ["p", "p"]
    assert "tbl" in blocks[:-3]
    paragraphs = _body_paragraphs(blob)
    assert paragraphs[-2].text == FAKE_FOOTER[0]
    assert paragraphs[-1].text == f"Дата: , ФИО: {FAKE_FIO}"
    assert all(p.alignment == WD_ALIGN_PARAGRAPH.CENTER for p in paragraphs[-2:])


def test_build_blank_document_without_footer_lines_prints_nothing():
    """Подвала в снимке нет (или он пустой/из пробелов) — ничего не печатаем."""
    for lines in (None, [], [" ", ""]):
        blob = build_blank_document(_context(blank_footer_lines=lines))
        assert _body_blocks(blob)[-2:] == ["tbl", "sectPr"]


def test_build_blank_document_substitutes_all_placeholders():
    """Плейсхолдеры контракта печатаются значениями (шапка, подвал, текст шага)."""
    context = _context(
        blank_header_html=FAKE_HEADER, blank_footer_lines=list(FAKE_FOOTER)
    )
    context["steps"] = [
        dict(
            context["steps"][0],
            title="Согласование {fio}",
            stage_lines=["Должность {position}", "Табельный {tab_num}"],
        )
    ]
    text = _docx_text(build_blank_document(context))
    for expected in (
        FAKE_FIO,                    # {fio}
        FAKE_POSITION,               # {position}
        FAKE_SERVICE,                # {department}
        FAKE_TAB_NUM,                # {tab_num}
        FAKE_ENTERPRISE_NAME,        # {enterprise}
        FAKE_BLANK,                  # {blank_name}
        FAKE_SUBJECT,                # {subject}
        FAKE_CONTENT,                # {content}
        "01.10.2026",                # {date}, формат ДД.ММ.ГГГГ
        "15.10.2026",                # {dismissal_date}
        "количество шагов: 1",       # {steps} по числу шагов заявки
        FAKE_MANAGER_FIO,            # {manager}
    ):
        assert expected in text, expected
    # Неподставленных плейсхолдеров на бланке не осталось.
    assert "{" not in text


def test_build_blank_document_unknown_placeholder_stays_text():
    """Неизвестный плейсхолдер остаётся текстом: печать не падает."""
    blob = build_blank_document(
        _context(
            blank_header_html="<p>{не_существует} и {unknown} для {fio}</p>",
            blank_footer_lines=["Подпись {steps_count}"],
        )
    )
    paragraphs = [paragraph.text for paragraph in _body_paragraphs(blob)]
    assert paragraphs[1] == f"{{не_существует}} и {{unknown}} для {FAKE_FIO}"
    assert paragraphs[-1] == "Подпись {steps_count}"


def test_build_blank_document_placeholders_without_dates_are_empty():
    """Нет даты/руководителя в контексте — плейсхолдеры печатаются пустыми."""
    blob = build_blank_document(
        _context(
            date=None,
            dismissal_date="",
            manager=None,
            blank_header_html="<p>Печать: [{date}] увольнение: [{dismissal_date}] "
            "руководитель: [{manager}]</p>",
        )
    )
    paragraphs = [paragraph.text for paragraph in _body_paragraphs(blob)]
    assert paragraphs[1] == "Печать: [] увольнение: [] руководитель: []"


def test_blank_placeholders_formats_dates_and_step_count():
    """Даты печатаются ДД.ММ.ГГГГ в любом разбираемом виде, {steps} — счётчик."""
    values = blank_placeholders(_context())
    assert values["date"] == "01.10.2026"
    assert values["dismissal_date"] == "15.10.2026"
    assert values["steps"] == "количество шагов: 2"
    assert values["manager"] == FAKE_MANAGER_FIO
    objects = blank_placeholders(
        _context(date=datetime(2026, 10, 1), dismissal_date=date(2026, 10, 15))
    )
    assert objects["date"] == "01.10.2026"
    assert objects["dismissal_date"] == "15.10.2026"
    # Мусор в дате или её отсутствие — пусто (печать не падает).
    assert blank_placeholders(_context(date="не дата", dismissal_date=None)) == {
        **values,
        "date": "",
        "dismissal_date": "",
    }


def test_build_blank_document_step_text_substitutes_placeholders():
    """Плейсхолдеры подставляются и в текст шага (название и пункты)."""
    context = _context()
    context["steps"] = [
        {
            "order": 1,
            "owner": FAKE_GROUP,
            "title": "Согласование {fio}",
            "stage_lines": [
                "Должность {position}",
                "Табельный {tab_num}",
                "Предприятие {enterprise}",
                "{не_существует}",
            ],
            "fio": "",
            "status": STEP_PENDING,
            "done_at": "",
        }
    ]
    assert _steps_rows(build_blank_document(context))[1][1].splitlines() == [
        f"Согласование {FAKE_FIO}",
        f"Должность {FAKE_POSITION}",
        f"Табельный {FAKE_TAB_NUM}",
        f"Предприятие {FAKE_ENTERPRISE_NAME}",
        "{не_существует}",
    ]


def test_build_blank_document_step_placeholder_value_stays_plain_text():
    """Значение {content}/{subject} в тексте шага печатается текстом: ни начертаний,
    ни лишних абзацев из значения, печать не падает.

    Порядок печати: в текст шага значение подставляется ЭКРАНИРОВАННЫМ (как в
    HTML-шапке), поэтому </p>/<b>/<script> из свободного текста заявки не
    разбираются как теги бланка — структура и начертания от значения не зависят.
    Разметка САМОГО текста шага (теги бланка) печатается как раньше."""
    markup = "</p><b>жирно</b> и <script>alert(1)</script>"
    context = _context(subject=f"тема {markup}", content=f"прощайте {markup}")
    context["steps"] = [
        {
            "order": 1,
            "owner": FAKE_GROUP,
            "title": "Этап <b>согласование</b> — {content}",
            "stage_lines": ["Пункт {subject}", "Пункт {не_существует}"],
            "fio": "",
            "status": STEP_PENDING,
            "done_at": "",
        }
    ]
    blob = build_blank_document(context)
    cell = _opened(blob).tables[-1].rows[1].cells[1]
    # Абзацев в ячейке ровно по одному на название и на пункт этапа.
    assert len(cell.paragraphs) == 3
    # Значение напечатано символами текста, неизвестный плейсхолдер остался текстом.
    assert [paragraph.text for paragraph in cell.paragraphs] == [
        f"Этап согласование — прощайте {markup}",
        f"Пункт тема {markup}",
        "Пункт {не_существует}",
    ]
    # Начертание — только у разметки самого шага (название этапа печатается
    # жирным), значение плейсхолдера начертания не дало.
    bold = [run.text for paragraph in cell.paragraphs for run in paragraph.runs if run.bold]
    assert bold == ["Этап ", "согласование"]
    # Документ собран целиком (значение печать не сломало): таблица шагов на месте.
    assert _docx_names(blob)[0] == "[Content_Types].xml"
    table = _opened(blob).tables[-1]
    assert table.rows[0].cells[0].text == "№"


# ---------------------------------------------------------------------------
# Генератор: DOCX из данных -> PDF (LibreOffice), без файлов-шаблонов
# ---------------------------------------------------------------------------


def test_generate_bypass_signature_has_no_template_arguments():
    """Пути через файлы-шаблоны в печати нет: ни body, ни .docx-файла."""
    params = inspect.signature(generate_bypass).parameters
    assert "template_body" not in params
    assert "template_file" not in params


def test_generate_bypass_builds_from_context_and_qr_url(tmp_path, monkeypatch):
    """В сборщик идёт контекст с qr_url = URL заявки (payload QR)."""
    captured = {}
    monkeypatch.setattr(
        docs_module, "build_blank_document",
        lambda context: captured.update(context) or b"docx-bytes",
    )
    monkeypatch.setattr(docs_module, "_render_qr_png", lambda url: b"qr-bytes")
    monkeypatch.setattr(docs_module, "_soffice_binary", lambda: None)
    result = generate_bypass(
        request_id="REQ-0001", version="v1", context=_context(),
        base_url=BASE_URL, files_dir=str(tmp_path),
    )
    assert captured["qr_url"] == QR_URL
    assert captured["blank_layout"] == "office"
    assert result.generated is False  # нет soffice — PDF не создан


def test_generate_bypass_offline_missing_library_reason(tmp_path, monkeypatch):
    """Нет библиотеки рендера → generated=False с причиной «офлайн»."""

    def missing(context):
        raise ModuleNotFoundError("No module named 'docx'")

    monkeypatch.setattr(docs_module, "build_blank_document", missing)
    result = generate_bypass(
        request_id="REQ-0001", version="v1", context=_context(),
        base_url=BASE_URL, files_dir=str(tmp_path),
    )
    assert result.generated is False
    assert result.reason.startswith("офлайн")
    assert "docx" in result.reason


# ---------------------------------------------------------------------------
# Ручка печати: макет из снимка в blank_kind, doc_templates не читается
# ---------------------------------------------------------------------------

def test_print_uses_blank_snapshot_and_reports_layout(
    client, requests_store, monkeypatch, tmp_path
):
    """Печать собирается из снимка бланка заявки; blank_kind = макет бланка."""
    requests_store.create(_request())
    captured = {}
    pdf = tmp_path / "bypass.pdf"
    pdf.write_bytes(b"%PDF-1.4 fake")

    def fake(**kwargs):
        captured.update(kwargs)
        return _fake_result(str(pdf))

    monkeypatch.setattr("app.documents.generate_bypass", fake)
    response = client.post("/requests/REQ-0001/print", headers=_hr_headers())
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["generated"] is True and body["blank_kind"] == "line"
    assert "template_body" not in captured and "template_file" not in captured
    context = captured["context"]
    assert context["blank_name"] == FAKE_BLANK
    assert context["blank_layout"] == "line"
    assert context["tab_num"] == FAKE_TAB_NUM
    assert context["enterprise"] == FAKE_ENTERPRISE_NAME  # названием из справочника
    assert context["steps"][0]["title"] == FAKE_STAGE
    assert context["steps"][0]["stage_lines"] == ["Проверить расчёты"]


def test_print_blank_kind_defaults_to_office(
    client, requests_store, monkeypatch, tmp_path
):
    """Без макета в снимке (старая заявка) печатаем офисный бланк."""
    requests_store.create(_request(blank_layout=None, blank_name=None, blank_version=None))
    pdf = tmp_path / "bypass.pdf"
    pdf.write_bytes(b"%PDF-1.4 fake")
    monkeypatch.setattr(
        "app.documents.generate_bypass", lambda **kwargs: _fake_result(str(pdf))
    )
    body = client.post("/requests/REQ-0001/print", headers=_hr_headers()).json()
    assert body["blank_kind"] == "office"


def test_print_ignores_doc_templates_setting(client, requests_store, monkeypatch, tmp_path):
    """Ключ doc_templates печать не читает: маркер шаблона в данных печати нет."""
    requests_store.create(_request())
    captured = {}
    pdf = tmp_path / "bypass.pdf"
    pdf.write_bytes(b"%PDF-1.4 fake")

    def fake(**kwargs):
        captured.update(kwargs)
        return _fake_result(str(pdf))

    monkeypatch.setattr("app.documents.generate_bypass", fake)
    response = client.post("/requests/REQ-0001/print", headers=_hr_headers())
    assert response.status_code == 200, response.text
    assert "ШАБЛОН-НЕ-ИСПОЛЬЗУЕТСЯ" not in str(captured)
    assert base64.b64decode(response.json()["pdf_b64"]) == b"%PDF-1.4 fake"


def test_print_passes_header_footer_snapshot_to_document(
    client, requests_store, monkeypatch, tmp_path
):
    """Печать берёт шапку/подвал и тему/содержание из снимка заявки."""
    requests_store.create(
        _request(
            blank_header_html=FAKE_HEADER,
            blank_footer_lines=list(FAKE_FOOTER),
            subject=FAKE_SUBJECT,
            content=FAKE_CONTENT,
        )
    )
    captured = {}
    pdf = tmp_path / "bypass.pdf"
    pdf.write_bytes(b"%PDF-1.4 fake")

    def fake(**kwargs):
        captured.update(kwargs)
        return _fake_result(str(pdf))

    monkeypatch.setattr("app.documents.generate_bypass", fake)
    response = client.post("/requests/REQ-0001/print", headers=_hr_headers())
    assert response.status_code == 200, response.text
    context = captured["context"]
    assert context["blank_header_html"] == FAKE_HEADER
    assert context["blank_footer_lines"] == list(FAKE_FOOTER)
    assert context["subject"] == FAKE_SUBJECT
    assert context["content"] == FAKE_CONTENT


def test_print_without_header_footer_snapshot(
    client, requests_store, monkeypatch, tmp_path
):
    """Заявка без шапки/подвала в снимке: печать получает пустые значения."""
    requests_store.create(_request())
    captured = {}
    pdf = tmp_path / "bypass.pdf"
    pdf.write_bytes(b"%PDF-1.4 fake")

    def fake(**kwargs):
        captured.update(kwargs)
        return _fake_result(str(pdf))

    monkeypatch.setattr("app.documents.generate_bypass", fake)
    response = client.post("/requests/REQ-0001/print", headers=_hr_headers())
    assert response.status_code == 200, response.text
    context = captured["context"]
    assert context["blank_header_html"] == ""
    assert context["blank_footer_lines"] == []


def test_print_422_without_steps(client, requests_store):
    """Шагов нет — 422 с понятным текстом (даже при заполненном doc_templates)."""
    requests_store.create(_request(steps=[]))
    response = client.post("/requests/REQ-0001/print", headers=_hr_headers())
    assert response.status_code == 422
    assert "шаг" in response.json()["detail"].lower()


def test_print_offline_is_not_500(client, requests_store, monkeypatch):
    """Офлайн/нет LibreOffice → 200 {"generated": false, reason}, не 500."""
    requests_store.create(_request())

    def missing(context):
        raise ModuleNotFoundError("No module named 'PIL'")

    monkeypatch.setattr(docs_module, "build_blank_document", missing)
    response = client.post("/requests/REQ-0001/print", headers=_hr_headers())
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["generated"] is False
    assert "PIL" in body["reason"]


def test_print_forbidden_for_step_owner(client, requests_store):
    """Печать — только разрешенной группе: владелец шага 403, аноним 401."""
    requests_store.create(_request())
    owner = {
        "X-Mock-Sam": "step.buhgalter",
        "X-Mock-Fio": _b64(FAKE_ASSIGNEE_FIO),
        "X-Mock-Mail": _b64("step.buhgalter@example.com"),
        "X-Mock-Department": _b64("Бухгалтерия"),
        "X-Mock-Title": _b64("Главный бухгалтер"),
        "X-Mock-Groups": FAKE_GROUP,
    }
    assert client.post("/requests/REQ-0001/print", headers=owner).status_code == 403
    assert client.post("/requests/REQ-0001/print", headers={}).status_code == 401


class _FakeEmployeeStore:
    """Зеркало сотрудников: одна строка по табелю (тестовая подмена)."""

    def __init__(self, rows=None, fail=False) -> None:
        self.rows = rows if rows is not None else [
            {
                "enterprise": "ENT",
                "tab_num": "Т-000777",
                "fio": "Вымышленнов Иван Тестович",
                "ad_sam": "vymyshlannyy",
                "dismissal_date": "2026-11-30",
            }
        ]
        self.fail = fail

    def search(self, enterprise, q, limit, offset=0):
        if self.fail:
            raise RuntimeError("зеркало сотрудников недоступно")
        return [dict(row) for row in self.rows]

    def find_by_request_key(self, enterprise, base_code, tab_num):
        # Боевой контракт: точное совпадение по предприятию/табелю, база
        # уточняет выборку, увольнение НЕ отсекается (печати нужна дата).
        if self.fail:
            raise RuntimeError("зеркало сотрудников недоступно")
        wanted = str(tab_num or "").strip().casefold()
        for row in self.rows:
            if str(row.get("tab_num") or "").strip().casefold() != wanted:
                continue
            if str(row.get("enterprise") or enterprise) != str(enterprise):
                continue
            if base_code and str(row.get("base_code") or "") != str(base_code):
                continue
            return dict(row)
        return None


class _FakeUserCards:
    """Зеркало AD users: карточки сотрудника и руководителя (тестовая подмена)."""

    def __init__(self, cards=None, managers=None) -> None:
        # None - значения по умолчанию, пустой словарь - «данных нет» (важно для теста).
        self.cards = {
            "vymyshlannyy": {
                "sam": "vymyshlannyy",
                "manager_dn": "CN=Ruk,OU=SED,DC=example,DC=local",
            },
            "Ruk": {"sam": "Ruk", "fio_full": "Вымышленнова Мария Тестовна", "manager_dn": None},
        } if cards is None else cards
        self.managers = (
            {"CN=Ruk,OU=SED,DC=example,DC=local": "Ruk"} if managers is None else managers
        )

    def user_card(self, sam):
        # Ключи в тесте смешанного регистра, поиск логина - без учёта регистра.
        wanted = str(sam or "").strip().casefold()
        for key, value in self.cards.items():
            if str(key).strip().casefold() == wanted:
                return dict(value)
        return None

    def manager_sam_by_dn(self, manager_dn):
        return self.managers.get(str(manager_dn or "").strip())


def test_hr_placeholders_filled_from_mirrors():
    """{date} — дата печати, {dismissal_date} — из кадровых данных, {manager} — ФИО руководителя."""
    from app.documents import _fill_hr_placeholders

    context = {}
    request = SimpleNamespace(enterprise="ENT", tab_num="Т-000777")
    _fill_hr_placeholders(context, request, _FakeUserCards(), _FakeEmployeeStore())

    values = blank_placeholders(context)
    assert values["dismissal_date"] == "30.11.2026"
    assert values["manager"] == "Вымышленнова Мария Тестовна"
    assert values["date"]  # дата печати всегда есть


def test_hr_placeholders_empty_when_mirror_unavailable():
    """Зеркала недоступны или сотрудник не найден — печать не падает, поля пустые."""
    from app.documents import _fill_hr_placeholders

    context = {}
    request = SimpleNamespace(enterprise="ENT", tab_num="Т-000777")

    _fill_hr_placeholders(context, request, None, _FakeEmployeeStore(fail=True))
    values = blank_placeholders(context)
    assert values["dismissal_date"] == "" and values["manager"] == ""

    context = {}
    _fill_hr_placeholders(context, request, _FakeUserCards(), _FakeEmployeeStore(rows=[]))
    values = blank_placeholders(context)
    assert values["dismissal_date"] == "" and values["manager"] == ""


def test_hr_placeholders_manager_absent_in_ad():
    """Руководителя в AD нет — {manager} пустой, печать не падает."""
    from app.documents import _fill_hr_placeholders

    context = {}
    request = SimpleNamespace(enterprise="ENT", tab_num="Т-000777")
    _fill_hr_placeholders(
        context, request, _FakeUserCards(managers={}), _FakeEmployeeStore()
    )
    values = blank_placeholders(context)
    assert values["manager"] == ""
    assert values["dismissal_date"] == "30.11.2026"


def test_employee_row_uses_request_key_lookup():
    """Строка сотрудника читается по ключу заявки (предприятие/база/табель).

    Точность совпадения обеспечивает сам запрос хранилища (проверяется в
    test_employee_hr_sync.py), здесь — отбор строки заглушкой."""
    from app.documents import _employee_row

    store = _FakeEmployeeStore(
        rows=[
            {"enterprise": "ENT", "tab_num": "Т-0007", "fio": "Похожий Тестов"},
            {
                "enterprise": "ENT",
                "base_code": "zup",
                "tab_num": "Т-000777",
                "fio": "Вымышленнов Иван Тестович",
            },
        ]
    )
    found = _employee_row(store, "ENT", "zup", "  Т-000777  ")
    assert found and found["fio"] == "Вымышленнов Иван Тестович"
    # База не совпала — это другой сотрудник, строки нет (печать не падает).
    assert _employee_row(store, "ENT", "drugaia", "Т-000777") is None


def test_render_blank_text_escapes_values_for_html_header():
    """В HTML-шапке значения плейсхолдеров печатаются текстом, а не разметкой.

    Свободный текст заявки ({content}) не должен разрывать абзац шапки и не
    должен превращаться в начертание."""
    context = {"content": "прощайте </p><b>жирно</b>", "fio": "Вымышленнов И.Т."}
    escaped = render_blank_text("{content} — {fio}", context, escape=True)
    assert "&lt;/p&gt;" in escaped and "<b>" not in escaped
    assert escaped == "прощайте &lt;/p&gt;&lt;b&gt;жирно&lt;/b&gt; — Вымышленнов И.Т."
    # Обычный текст (подвал, шаги) подставляется как есть — без разметки.
    plain = render_blank_text("{content}", context)
    assert plain == context["content"]

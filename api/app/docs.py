# Интерфейсы и offline-реализация генерации бланков (волна B3).
# Генерация — логически часть worker, не api-запроса (см. скил mail-docs).
# Здесь только stdlib: DOCX собирается вручную (zip + document.xml),
# QR — заглушка (PNG с URL заявки в tEXt-блоке, матрица — на стенде
# библиотекой qrcode).
# Печать бланков (волна «Справочник бланков»): документ собирается ИЗ ДАННЫХ
# заявки — build_blank_document на python-docx, макет из снимка blank_layout
# (office|line). Файлы-шаблоны .docx источником оформления больше не являются
# (ключ настроек doc_templates и ручки файлов-шаблонов удалены — печать собирает
# бланок из данных).
# ПДн (остаток отпуска и др.) в бланки не включаются — см. sanitize_context.
# Зависимости стенда (python-docx/qrcode/Pillow/LibreOffice) — раскомментирует
# стенд, offline только stdlib.

from __future__ import annotations

import io
import logging
import re
import struct
import zipfile
import zlib
from dataclasses import dataclass
from datetime import datetime
from html import unescape
from pathlib import Path
from typing import Dict, List, Protocol

# Логгер модуля: сбои генерации бегунка пишем в журнал (см. generate_bypass).
_LOGGER = logging.getLogger(__name__)

# Метка PDF-стаба: рендер настоящего PDF — только на стенде через LibreOffice.
PDF_STUB_MARK = "рендер на стенде (LibreOffice)"

# Поля, запрещенные в бланках/письмах (ПДн, только для ОК).
FORBIDDEN_CONTEXT_FIELDS = frozenset({"vacation_balance", "vacation", "mail", "manager_dn"})

# Подстановка вида {{ имя }} (минимальный Jinja для offline; на стенде — настоящий Jinja).
_PLACEHOLDER_RE = re.compile(r"\{\{\s*(\w+)\s*\}\}")


def sanitize_context(context: Dict[str, object]) -> Dict[str, object]:
    """Убрать из контекста ПДн-поля (в бланки/письма они не включаются)."""
    return {k: v for k, v in context.items() if k not in FORBIDDEN_CONTEXT_FIELDS}


def render_text(template: str, context: Dict[str, object]) -> str:
    """Подставить {{var}} из контекста; неизвестное — пустая строка."""
    safe = sanitize_context(context)
    return _PLACEHOLDER_RE.sub(lambda m: str(safe.get(m.group(1), "")), template)


def request_url(base_url: str, request_id: str) -> str:
    """Ссылка на заявку для QR (базовый URL — из настроек, не из кода)."""
    return base_url.rstrip("/") + "/requests/" + request_id


def next_version(existing: List[int]) -> int:
    """Следующая версия бланка (v1/v2...); старые версии не удаляются."""
    return (max(existing) if existing else 0) + 1


def qr_stub_png(url: str) -> bytes:
    """Заглушка QR: валидный PNG, URL заявки — в tEXt-блоке (читается тестом).

    Настоящая QR-матрица генерируется на стенде библиотекой qrcode;
    здесь — белый квадрат 8x8 + текстовый чанк с URL (stdlib: struct+zlib).
    """
    url_bytes = url.encode("utf-8")

    def chunk(ctype: bytes, data: bytes) -> bytes:
        # Длина + тип + данные + CRC32 (формат PNG).
        return struct.pack(">I", len(data)) + ctype + data + struct.pack(">I", zlib.crc32(ctype + data))

    sig = b"\x89PNG\r\n\x1a\n"
    ihdr = struct.pack(">IIBBBBB", 8, 8, 8, 2, 0, 0, 0)  # 8x8 truecolor
    # tEXt: ключевое слово + нулевой байт + текст (URL заявки).
    text = b"RequestURL\x00" + url_bytes
    # Белый квадрат 8x8 (каждая строка: фильтр-байт + пиксели).
    raw = b"".join(b"\x00" + b"\xff\xff\xff" * 8 for _ in range(8))
    return sig + chunk(b"IHDR", ihdr) + chunk(b"tEXt", text) + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b"")


def _xml_escape(text: str) -> str:
    """Экранирование для document.xml."""
    return (
        text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")
    )


def _paragraph(text: str) -> str:
    """Один абзац WordprocessingML (каждая строка — свой абзац)."""
    return "<w:p><w:r><w:t xml:space=\"preserve\">" + _xml_escape(text) + "</w:t></w:r></w:p>"


@dataclass(frozen=True)
class RenderedDoc:
    """Результат рендера бланка (байты + метаданные версии)."""

    request_id: str
    version: int
    url: str
    docx: bytes
    pdf_stub: bytes


class DocRenderer(Protocol):
    """Граница рендера бланков (на стенде — реализация на python-docx-template)."""

    def render_docx(
        self,
        request_id: str,
        version: int,
        template_body: str,
        context: Dict[str, object],
        base_url: str,
    ) -> bytes:
        """Собрать DOCX-бегунок из переданного текста-шаблона (Jinja-подстановка)."""
        ...  # pragma: no cover

    def render_pdf_stub(self, request_id: str, version: int) -> bytes:
        """Стаб PDF с пометкой про LibreOffice (настоящий рендер — на стенде)."""
        ...  # pragma: no cover


class StdlibDocxRenderer:
    """Offline-рендер DOCX на stdlib (zipfile): текст + QR-заглушка + версия.

    Остался от волны B3 как заглушка офлайна; печать заявки (documents.py) его
    не вызывает — она собирает бланк из данных через build_blank_document.
    """

    def render_docx(
        self,
        request_id: str,
        version: int,
        template_body: str,
        context: Dict[str, object],
        base_url: str,
    ) -> bytes:
        """Собрать DOCX: отрендеренный текст + request_id + версия + URL заявки."""
        url = request_url(base_url, request_id)
        body = render_text(template_body, context)
        lines = [body, "Заявка: " + request_id, "Версия: v" + str(version), "QR: " + url]
        paragraphs = "".join(_paragraph(line) for line in lines)
        document_xml = (
            "<?xml version=\"1.0\" encoding=\"UTF-8\" standalone=\"yes\"?>"
            "<w:document xmlns:w=\"http://schemas.openxmlformats.org/wordprocessingml/2006/main\">"
            "<w:body>" + paragraphs + "</w:body></w:document>"
        )
        content_types = (
            "<?xml version=\"1.0\" encoding=\"UTF-8\"?>"
            "<Types xmlns=\"http://schemas.openxmlformats.org/package/2006/content-types\">"
            "<Default Extension=\"rels\" ContentType=\"application/vnd.openxmlformats-package.relationships+xml\"/>"
            "<Default Extension=\"xml\" ContentType=\"application/xml\"/>"
            "<Default Extension=\"png\" ContentType=\"image/png\"/>"
            "<Override PartName=\"/word/document.xml\" "
            "ContentType=\"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml\"/>"
            "</Types>"
        )
        rels = (
            "<?xml version=\"1.0\" encoding=\"UTF-8\"?>"
            "<Relationships xmlns=\"http://schemas.openxmlformats.org/package/2006/relationships\">"
            "<Relationship Id=\"rId1\" "
            "Type=\"http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument\" "
            "Target=\"word/document.xml\"/></Relationships>"
        )
        # Связь картинки QR с документом (заглушка; на стенде — настоящая вставка).
        doc_rels = (
            "<?xml version=\"1.0\" encoding=\"UTF-8\"?>"
            "<Relationships xmlns=\"http://schemas.openxmlformats.org/package/2006/relationships\">"
            "<Relationship Id=\"rIdQr\" "
            "Type=\"http://schemas.openxmlformats.org/officeDocument/2006/relationships/image\" "
            "Target=\"media/qr.png\"/></Relationships>"
        )
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
            zf.writestr("[Content_Types].xml", content_types)
            zf.writestr("_rels/.rels", rels)
            zf.writestr("word/document.xml", document_xml)
            zf.writestr("word/_rels/document.xml.rels", doc_rels)
            zf.writestr("word/media/qr.png", qr_stub_png(url))
        return buf.getvalue()

    def render_pdf_stub(self, request_id: str, version: int) -> bytes:
        """Минимальный PDF-стаб с пометкой про стенд (не настоящий рендер)."""
        # Минимальный однострочный PDF: текст с пометкой и номером заявки/версией.
        text = "Bypass %s v%d (%s)" % (request_id, version, PDF_STUB_MARK)
        # Экранирование скобок по спецификации PDF (иначе строка рвется).
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        content = "BT /F1 12 Tf 72 720 Td (" + escaped + ") Tj ET"
        content_bytes = content.encode("utf-8")
        header = b"%PDF-1.4\n"
        # Комментарий с пометкой дословно (проверяется тестом; комментарии в PDF легальны).
        comment = ("% " + PDF_STUB_MARK + "\n").encode("utf-8")
        header = header + comment
        body = (
            b"1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n"
            b"2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj\n"
            b"3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 595 842]/Contents 4 0 R>>endobj\n"
            + b"4 0 obj<</Length "
            + str(len(content_bytes)).encode("ascii")
            + b">>stream\n"
            + content_bytes
            + b"\nendstream\nendobj\n"
        )
        trailer = b"trailer<</Root 1 0 R>>\n%%EOF"
        return header + body + trailer

    def render(
        self,
        request_id: str,
        version: int,
        template_body: str,
        context: Dict[str, object],
        base_url: str,
    ) -> RenderedDoc:
        """Удобная обертка: DOCX + PDF-стаб одним вызовом."""
        url = request_url(base_url, request_id)
        return RenderedDoc(
            request_id=request_id,
            version=version,
            url=url,
            docx=self.render_docx(request_id, version, template_body, context, base_url),
            pdf_stub=self.render_pdf_stub(request_id, version),
        )


# ---------------------------------------------------------------------------
# W3a: боевая генерация бегунка (DOCX -> PDF + QR) на стенде.
# Зависимости стенда (python-docx/qrcode/Pillow/LibreOffice)
# импортируются лениво: офлайн их нет — generate_bypass вернет generated=False
# с причиной (не 500); прочие сбои не глотаются. Оформление бланка — из данных
# заявки и снимка бланка (см. build_blank_document), файлов-шаблонов нет.
# Профиль soffice — в /tmp (см. _convert_to_pdf_report): HOME контейнера не
# существует, софт без профиля RC=77.
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class BypassResult:
    """Итог генерации бегунка: файлы на диске или причина отказа (не ошибка)."""

    generated: bool = False
    reason: str = ""
    docx_path: str = ""
    pdf_path: str = ""
    qr_payload: str = ""


def _step_done_text(value: object) -> str:
    """Дата отметки шага строкой ДД.ММ.ГГГГ ЧЧ:ММ (пусто — шаг не отмечен).

    Принимаем и datetime модели шага, и строку (уже отформатированную дату) —
    контекст собирается из «птичьего» объекта запроса, как и весь остальной."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y %H:%M")
    text = str(value).strip()
    if not text:
        return ""
    try:
        return datetime.fromisoformat(text).strftime("%d.%m.%Y %H:%M")
    except ValueError:
        return text[:16].replace("T", " ")


def _step_logins(step: object) -> List[str]:
    """Логины ответственных шага из снимка (assignees, миграция 0013).

    Пустой снимок — прежний одиночный assignee (заявки, выданные до миграции):
    печатать и уведомлять надо прежнего исполнителя. Мусор в jsonb и дубли
    отсекаются, порядок снимка сохраняется."""
    logins: List[str] = []
    for raw in getattr(step, "assignees", None) or []:
        sam = str(raw).strip()
        if sam and sam not in logins:
            logins.append(sam)
    if not logins:
        single = str(getattr(step, "assignee", None) or "").strip()
        if single:
            logins.append(single)
    return logins


def build_bypass_context(request: object) -> Dict[str, object]:
    """Контекст бланка: поля 1С заявки без ПДн (mail/отпуск — не включаем)
    + снимок бланка (название/версия/макет) + шаги маршрута (владельцы шагов —
    участники процесса, не ПДн).

    По каждому шагу отдаём снимок этапа из справочников (title/stage_lines),
    снимок ответственных (assignees — все логины, один ответственный — список
    из одного) и отметку (done_at): бланк печатается по маршруту, собранному из
    этапов, и в колонке «Ответственный» печатаются ВСЕ ответственные шага.
    assignee_names здесь пустой намеренно: ФИО по логинам подставляет
    documents.py из зеркала AD (docs.py про хранилище заявок не знает — импорт
    был бы циклическим, его не делаем); fio — ФИО первого ответственного
    (прежнее поле контекста, оставлено для одиночного исполнителя)."""
    steps = []
    for step in sorted(request.steps, key=lambda s: s.order):
        owner = getattr(step, "assignee", None) or step.owner_group
        stage_lines = [
            str(line)
            for line in (getattr(step, "stage_lines", None) or [])
            if str(line).strip()
        ]
        steps.append(
            {
                "order": step.order,
                "owner": owner,
                "status": step.status,
                # Название этапа (снимок этапа) и текст его пунктов: в печать
                # идут через build_blank_document (HTML из редактора — там).
                "title": str(
                    getattr(step, "stage_title", None) or step.owner_group or ""
                ),
                "position": str(getattr(step, "stage_title", None) or step.owner_group or ""),
                # Снимок ответственных: логины (assignees, миграция 0013) плюс
                # прежний одиночный assignee — у заявок, выданных до миграции,
                # снимка нет, печатать надо прежнего исполнителя.
                "assignees": _step_logins(step),
                "assignee": getattr(step, "assignee", None) or "",
                "fio": "",
                # ФИО ответственных по логинам: подставляет documents.py
                # (в бланке колонка «Ответственный» — по строке на ответственного).
                "assignee_names": [],
                "stage_lines": stage_lines or [owner],
                "done_at": _step_done_text(getattr(step, "done_at", None)),
            }
        )
    return {
        "request_id": request.id,
        "fio": request.fio,
        "department": request.department,
        "position": request.position,
        "category": request.category or "",
        "enterprise": request.enterprise or "",
        "tab_num": getattr(request, "tab_num", None) or "",
        "doc_type_code": getattr(request, "doc_type_code", None) or "",
        # Снимок выбранного бланка: правка справочника не меняет выданную
        # заявку, поэтому печатаем ровно тот макет, который зафиксирован.
        "blank_id": getattr(request, "blank_id", None),
        "blank_name": getattr(request, "blank_name", None) or "",
        "blank_version": getattr(request, "blank_version", None),
        "blank_layout": getattr(request, "blank_layout", None) or "",
        "steps": steps,
    }


# ---------------------------------------------------------------------------
# Бланк из данных (волна «Справочник бланков»): DOCX собирается python-docx,
# макет — из снимка заявки (blank_layout), два встроенных пресета. Файлы-шаблоны
# .docx печать не читает. Зависимости стенда (python-docx/qrcode/Pillow)
# импортируются лениво: офлайн их нет — generate_bypass вернет generated=False
# с причиной (не 500).
# ---------------------------------------------------------------------------

# Шрифт печатной формы бланка (константа вёрстки, не настройка).
BLANK_FONT_NAME = "Times New Roman"

# Два встроенных пресета оформления; выбирает blank_layout из снимка заявки.
# Числа — константы вёрстки печатной формы: поля страницы (В/П/Н/Л, мм), кегль
# обычного текста и заголовка, сторона QR (мм), сетка полей шапки, заливка
# шапки таблицы, доли ширины колонок таблицы шагов.
LAYOUT_PRESETS: Dict[str, dict] = {
    "office": {
        "margins_mm": (20, 15, 18, 15),
        "font_pt": 11,
        "title_pt": 14,
        "qr_mm": 30,
        "grid_info": True,
        "shade_header": True,
        "step_widths": (0.07, 0.55, 0.26, 0.12),
    },
    "line": {
        "margins_mm": (12, 10, 12, 10),
        "font_pt": 9,
        "title_pt": 11,
        "qr_mm": 20,
        "grid_info": False,
        "shade_header": False,
        "step_widths": (0.06, 0.56, 0.26, 0.12),
    },
}

# Макет по умолчанию, если в снимке пусто или значение не из списка пресетов
# (дефолт совпадает с default='office' миграции 0012).
DEFAULT_BLANK_LAYOUT = "office"

# Подписи полей шапки бланка: (подпись, ключ контекста).
_HEADER_FIELDS = (
    ("ФИО", "fio"),
    ("Должность", "position"),
    ("Служба", "department"),
    ("Табельный номер", "tab_num"),
    ("Предприятие", "enterprise"),
)

# Колонки таблицы шагов.
_STEP_HEADERS = ("№", "Этап", "Ответственный", "Отметка/дата")

# Теги текста этапа, разрешённые к печати (визуальный редактор отдаёт HTML):
# начертание и перенос строки. Остальные теги — обычный текст без разметки,
# содержимое служебных тегов на бланок не попадает вовсе.
_BOLD_TAGS = frozenset({"b", "strong"})
_ITALIC_TAGS = frozenset({"i", "em"})
_UNDERLINE_TAGS = frozenset({"u"})
_BREAK_TAGS = frozenset({"br", "p", "li"})
_HIDDEN_TAGS = frozenset(
    {
        "script", "style", "head", "title", "noscript", "template",
        "iframe", "object", "embed", "svg", "math",
    }
)
_HTML_TAG_RE = re.compile(r"<\s*(/?)\s*([a-zA-Z][a-zA-Z0-9]*)[^>]*>")


def resolve_blank_layout(value: object) -> str:
    """Макет бланка по значению снимка: office|line; всё прочее (пусто, мусор) —
    office. Единая точка выбора макета для печати и ответа API (blank_kind)."""
    name = str(value or "").strip().lower()
    return name if name in LAYOUT_PRESETS else DEFAULT_BLANK_LAYOUT


def _blank_title(context: Dict[str, object]) -> str:
    """Название бланка на бланке: снимок blank_name, иначе вид документа."""
    return (
        str(context.get("blank_name") or "").strip()
        or str(context.get("doc_type_code") or "").strip()
        or "Бланк"
    )


def _printable_width_mm(document: object) -> float:
    """Ширина наборной полосы страницы, мм (поля из пресета)."""
    from docx.shared import Emu

    section = document.sections[0]
    # Разность длин Emu — обычный int, в миллиметры возвращает сам Emu.
    return Emu(section.page_width - section.left_margin - section.right_margin).mm


def _docx_document(preset: dict):
    """Пустой документ A4 под макет: поля страницы и базовый шрифт."""
    from docx import Document
    from docx.shared import Mm, Pt

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    top, right, bottom, left = preset["margins_mm"]
    section.top_margin, section.right_margin = Mm(top), Mm(right)
    section.bottom_margin, section.left_margin = Mm(bottom), Mm(left)
    style = document.styles["Normal"]
    style.font.name = BLANK_FONT_NAME
    style.font.size = Pt(preset["font_pt"])
    return document


def _shade_cell(cell: object, color: str = "D9D9D9") -> None:
    """Заливка ячейки (шапка таблицы офисного пресета)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shading)


def _add_runs(
    paragraph: object,
    text: str,
    *,
    bold: bool = False,
    italic: bool = False,
    underline: bool = False,
) -> None:
    """Текст в абзац с начертанием; переносы строк — разрывом прогона."""
    for index, line in enumerate(text.split("\n")):
        if index:
            paragraph.add_run().add_break()
        if not line:
            continue
        run = paragraph.add_run(line)
        run.bold = bold
        run.italic = italic
        run.underline = underline


def _add_stage_text(paragraph: object, text: str, *, bold: bool = False) -> None:
    """Текст этапа (может прийти HTML из визуального редактора) в абзац DOCX.

    В печать идут только разрешённые теги (_BOLD_TAGS/_ITALIC_TAGS/
    _UNDERLINE_TAGS/_BREAK_TAGS); содержимое служебных тегов выбрасывается,
    любые другие теги идут обычным текстом без разметки. Свой санитайзер не
    пишем: ПДн-поля вычищает sanitize_context, здесь только отбор тегов,
    разрешённых к печати (контракт фазы).
    """
    cur_bold, cur_italic, cur_underline = bold, False, False
    hidden: List[str] = []
    emitted = False
    pos = 0
    for match in _HTML_TAG_RE.finditer(text):
        closing, tag = match.group(1) == "/", match.group(2).lower()
        chunk, pos = text[pos : match.start()], match.end()
        # Текст печатается, пока не открыт служебный тег: внутри него (script,
        # style и т.п.) на бланок не попадает ничего.
        if chunk and not hidden:
            _add_runs(
                paragraph,
                unescape(chunk),
                bold=cur_bold,
                italic=cur_italic,
                underline=cur_underline,
            )
            emitted = True
        if tag in _HIDDEN_TAGS:
            if hidden:
                if closing and hidden[-1] == tag:
                    hidden.pop()
            elif not closing:
                hidden.append(tag)
            continue
        if hidden:
            continue
        if tag in _BREAK_TAGS:
            if not closing and emitted:
                paragraph.add_run().add_break()
                emitted = True
        elif tag in _BOLD_TAGS:
            cur_bold = not closing
        elif tag in _ITALIC_TAGS:
            cur_italic = not closing
        elif tag in _UNDERLINE_TAGS:
            cur_underline = not closing
    tail = text[pos:]
    if tail and not hidden:
        _add_runs(
            paragraph,
            unescape(tail),
            bold=cur_bold,
            italic=cur_italic,
            underline=cur_underline,
        )


def _add_blank_header(document: object, context: Dict[str, object], preset: dict) -> None:
    """Шапка бланка: название, поля сотрудника, QR на заявку (qr_url контекста)."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Mm, Pt

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading = title.add_run(_blank_title(context))
    heading.bold = True
    heading.font.size = Pt(preset["title_pt"])
    values = [
        (label, str(context.get(key) or "").strip()) for label, key in _HEADER_FIELDS
    ]
    if preset["grid_info"]:
        # Офисный пресет: поля шапки сеткой (таблица без границ).
        grid = document.add_table(rows=len(values), cols=2)
        label_mm = 35.0
        width_mm = _printable_width_mm(document)
        for row, (label, value) in zip(grid.rows, values):
            row.cells[0].width = Mm(label_mm)
            row.cells[1].width = Mm(max(10.0, width_mm - label_mm))
            cell = row.cells[0].paragraphs[0]
            cell.add_run(label + ": ").bold = True
            row.cells[1].paragraphs[0].add_run(value)
    else:
        # Линейный пресет: поля шапки списком, компактнее.
        for label, value in values:
            paragraph = document.add_paragraph()
            paragraph.add_run(label + ": ").bold = True
            paragraph.add_run(value)
    qr_url = str(context.get("qr_url") or "").strip()
    if not qr_url:
        return
    picture = document.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    picture.add_run().add_picture(
        io.BytesIO(_render_qr_png(qr_url)),
        width=Mm(preset["qr_mm"]),
        height=Mm(preset["qr_mm"]),
    )
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    note = caption.add_run(qr_url)
    note.font.size = Pt(max(6, preset["font_pt"] - 3))


def _step_responsible(step: Dict[str, object]) -> List[str]:
    """Ответственные шага на бланке — список строк.

    ФИО всех персональных исполнителей снимка (assignee_names — их подставляет
    documents._fill_step_fio по логинам), по одному в строке: колонка узкая,
    а отметку/подпись каждый ставит напротив своей фамилии. Прежнее одиночное
    fio — как есть. Пустой снимок (групповой этап) — группа AD печатается как
    есть, на бумаге её пишут от руки."""
    names = [
        str(name).strip() for name in (step.get("assignee_names") or []) if str(name).strip()
    ]
    if names:
        return names
    single = str(step.get("fio") or "").strip()
    return [single or str(step.get("owner") or "").strip()]


def _step_marks(step: Dict[str, object]) -> list[str]:
    """Отметка/дата шага: статус и дата выполнения; без даты — место под росчерк."""
    done_at = str(step.get("done_at") or "").strip()
    return [
        line
        for line in (str(step.get("status") or "").strip(), done_at or "_" * 10)
        if line
    ]


def _add_steps_table(
    document: object, context: Dict[str, object], preset: dict
) -> None:
    """Таблица шагов: № | этап с его пунктами | ответственные | отметка/дата.

    В колонке «Ответственный» печатаются ВСЕ ответственные шага (реестровый этап
    с несколькими исполнителями), по одному в строке."""
    from docx.shared import Mm

    steps = [step for step in (context.get("steps") or []) if isinstance(step, dict)]
    table = document.add_table(rows=1, cols=len(_STEP_HEADERS))
    if preset["shade_header"]:
        table.style = "Table Grid"
    table.autofit = False
    width_mm = _printable_width_mm(document)
    widths = [Mm(width_mm * share) for share in preset["step_widths"]]
    for index, column in enumerate(table.columns):
        column.width = widths[index]
    for index, cell in enumerate(table.rows[0].cells):
        cell.width = widths[index]
        cell.paragraphs[0].add_run(_STEP_HEADERS[index]).bold = True
        if preset["shade_header"]:
            _shade_cell(cell)
    for position, step in enumerate(steps, start=1):
        row = table.add_row()
        for index, cell in enumerate(row.cells):
            cell.width = widths[index]
        row.cells[0].paragraphs[0].add_run(str(step.get("order") or position))
        stage = row.cells[1].paragraphs[0]
        _add_stage_text(stage, str(step.get("title") or "").strip(), bold=True)
        for line in step.get("stage_lines") or []:
            text = str(line).strip()
            if text:
                _add_stage_text(row.cells[1].add_paragraph(), text)
        responsible_cell = row.cells[2]
        for index, name in enumerate(_step_responsible(step)):
            paragraph = (
                responsible_cell.paragraphs[0] if index == 0
                else responsible_cell.add_paragraph()
            )
            paragraph.add_run(name)
        mark_cell = row.cells[3]
        for index, line in enumerate(_step_marks(step)):
            paragraph = (
                mark_cell.paragraphs[0] if index == 0 else mark_cell.add_paragraph()
            )
            paragraph.add_run(line)


def _add_blank_footer(
    document: object, context: Dict[str, object], preset: dict
) -> None:
    """Подвал бланка: номер заявки, название бланка и его версия (снимок)."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    version = context.get("blank_version")
    parts = [
        str(context.get("request_id") or "").strip(),
        _blank_title(context) + (f" v{version}" if version else ""),
    ]
    footer = document.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run(" · ".join(part for part in parts if part))


def build_blank_document(context: Dict[str, object]) -> bytes:
    """Бланк заявки как DOCX: собирается из данных (python-docx), не из файла.

    Макет — из снимка blank_layout (office|line, дефолт office, см.
    resolve_blank_layout), шапка берёт blank_name и поля сотрудника, шаги
    печатаются таблицей с ответственными — ВСЕМИ по снимку (ФИО подставил
    documents.py по логинам assignee_names, по строке на ответственного), группы
    AD — как есть, QR — на qr_url из контекста (его кладёт
    generate_bypass: URL заявки). ПДн-поля вычищает sanitize_context.
    Заявку без шагов не печатает пустым документом — печать отвечает 422
    (documents.print_bypass). Ожидает python-docx/qrcode/Pillow (стенд);
    офлайн ImportError разбирает generate_bypass (generated=False).
    """
    safe = sanitize_context(dict(context))
    preset = LAYOUT_PRESETS[resolve_blank_layout(safe.get("blank_layout"))]
    document = _docx_document(preset)
    _add_blank_header(document, safe, preset)
    _add_steps_table(document, safe, preset)
    _add_blank_footer(document, safe, preset)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _render_qr_png(url: str) -> bytes:
    """QR-матрица библиотекой qrcode (payload — URL заявки)."""
    import io

    import qrcode

    img = qrcode.make(url)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _soffice_binary() -> str | None:
    """Путь к soffice (системный LibreOffice на ВМ) либо None."""
    import shutil

    return shutil.which("soffice")


# Сколько последних строк stderr soffice показываем в причине отказа.
_SOFFICE_STDERR_LINES = 3


@dataclass(frozen=True)
class SofficeReport:
    """Итог конвертации: путь к PDF (либо None) + служебная диагностика."""

    pdf_path: str | None = None
    detail: str = ""


def _soffice_detail(returncode: int, stderr: str) -> str:
    """Код возврата soffice и хвост stderr — только служебный вывод конвертера
    (пути файлов, коды), ПДн заявки туда не попадают. Хвост ограничен, чтобы
    ответ API не раздувался шумом вроде сообщений dconf/fontconfig."""
    tail = [line for line in (stderr or "").splitlines() if line.strip()]
    detail = f"код возврата {returncode}"
    if tail:
        detail += ": " + " | ".join(line.strip()[:200] for line in tail[-_SOFFICE_STDERR_LINES:])
    return detail


def _convert_to_pdf_report(docx_path: str, out_dir: str, soffice: str | None) -> SofficeReport:
    """soffice --convert-to pdf с профилем в /tmp; PDF либо None + диагностика."""
    if soffice is None:
        return SofficeReport(detail="soffice не найден в PATH (LibreOffice не установлен)")
    import os
    import subprocess

    # Профиль LibreOffice — отдельный каталог в /tmp с суффиксом pid: HOME
    # контейнера (/home/appuser) не существует, и soffice без профиля падает с
    # RC=77 («application cannot be started»); общий профиль на все процессы
    # при параллельной печати даёт «another instance running».
    profile = f"-env:UserInstallation=file:///tmp/soffice-profile-{os.getpid()}"
    result = subprocess.run(
        [soffice, profile, "--headless", "--convert-to", "pdf", "--outdir", out_dir, docx_path],
        capture_output=True,
        text=True,
        timeout=60,
    )
    if result.returncode != 0:
        return SofficeReport(detail=_soffice_detail(result.returncode, result.stderr))
    pdf = str(Path(docx_path).with_suffix(".pdf"))
    if not Path(pdf).exists():
        return SofficeReport(detail=_soffice_detail(result.returncode, "PDF-файл не появился"))
    return SofficeReport(pdf_path=pdf)


def _convert_to_pdf(docx_path: str, out_dir: str, soffice: str | None) -> str | None:
    """soffice --convert-to pdf; путь к PDF либо None при сбое (без диагностики)."""
    return _convert_to_pdf_report(docx_path, out_dir, soffice).pdf_path


# Текст «No module named '...'» в ImportError/ModuleNotFoundError (у вручную
# созданных исключений атрибут name пустой — берём имя из текста).
_MISSING_MODULE_RE = re.compile(r"No module named [\"']([\w.]+)[\"']")


def _missing_library(exc: BaseException) -> str:
    """Имя модуля, которого нет: exc.name, иначе разбор текста ошибки.

    Для вложенных импортов берётся верхний уровень (PIL.Image -> PIL)."""
    name = getattr(exc, "name", None) or ""
    if not name:
        match = _MISSING_MODULE_RE.search(str(exc))
        name = match.group(1) if match else "неизвестно"
    return name.split(".")[0]


def generate_bypass(
    request_id: str,
    version: str,
    context: Dict[str, object],
    base_url: str,
    files_dir: str,
) -> BypassResult:
    """Собрать бланк заявки: DOCX (python-docx) -> PDF (LibreOffice) + QR.

    Оформление целиком из данных: build_blank_document по контексту (снимок
    blank_layout, шапка, таблица шагов со всеми ответственными); файлы-шаблоны
    .docx и ключ настроек doc_templates удалены. qr_url кладёт сам
    generate_bypass — payload QR и подпись под картинкой это URL заявки;
    sanitize_context применяется ДО сборки (qr_url/steps под фильтр ПДн
    не подпадают).

    Офлайн (нет python-docx/qrcode/PIL) или нет soffice —
    BypassResult(generated=False, reason=...): файлы не записываются в документы
    (в БД только мета успешной генерации). QR payload — URL заявки.
    Причина сбоя конвертации — код возврата soffice и хвост его stderr.
    Прочие сбои не глотаются: пишем в лог и пробрасываем наружу.
    """
    url = request_url(base_url, request_id)
    safe = sanitize_context(dict(context, request_id=request_id, qr_url=url))
    try:
        docx_bytes = build_blank_document(safe)
        qr_bytes = _render_qr_png(url)
    except ImportError as exc:
        # Офлайн: нет python-docx/qrcode/PIL (в т.ч. вложенный импорт).
        # Называем саму библиотеку, а не только класс исключения.
        return BypassResult(
            generated=False,
            reason=(
                f"офлайн: нет библиотеки {_missing_library(exc)} "
                f"({exc.__class__.__name__}: {exc})"
            ),
        )
    except Exception:
        # Настоящий сбой (сборка бланка, QR, файлы) — под «офлайн» его прятать нельзя.
        _LOGGER.exception("Ошибка генерации бегунка по заявке %s", request_id)
        raise
    out_dir = Path(files_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    docx_path = out_dir / f"bypass_{request_id}_v{version}.docx"
    docx_path.write_bytes(docx_bytes)
    report = _convert_to_pdf_report(str(docx_path), str(out_dir), _soffice_binary())
    pdf_path = report.pdf_path
    if pdf_path is None:
        # PDF не создан: бегунок не состоялся, файл не оставляем.
        try:
            docx_path.unlink(missing_ok=True)
        except OSError:
            pass
        return BypassResult(
            generated=False,
            reason=f"LibreOffice: PDF не создан ({report.detail})",
        )
    qr_path = out_dir / f"bypass_{request_id}_v{version}_qr.png"
    qr_path.write_bytes(qr_bytes)
    return BypassResult(
        generated=True,
        docx_path=str(docx_path),
        pdf_path=str(pdf_path),
        qr_payload=url,
    )
